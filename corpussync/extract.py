"""Read supported files into plain text. Missing pdf/docx extras skip the file."""

from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

from corpussync.vtt import clean_srt, clean_vtt

SUPPORTED_EXTENSIONS = {
    ".md",
    ".markdown",
    ".txt",
    ".vtt",
    ".srt",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
}

_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")


def _missing_extra(path: Path, extra: str) -> None:
    print(f'skip {path}: install the {extra} extra (pip install "corpussync[{extra}]")')


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _markdown_title(text: str) -> str | None:
    for line in text.splitlines():
        match = _HEADING.match(line)
        if match:
            title = match.group(1).strip()
            if title:
                return title
    return None


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title: str | None = None
        self._skip: str | None = None
        self._heading: str | None = None
        self._heading_parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "title"):
            self._skip = tag
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6") and self.title is None and self._skip is None:
            self._heading = tag
            self._heading_parts = []

    def handle_endtag(self, tag):
        if self._skip == tag:
            self._skip = None
        if self._heading == tag:
            title = " ".join(part for part in self._heading_parts if part).strip()
            if title:
                self.title = title
            self._heading = None

    def handle_data(self, data):
        if self._skip:
            return
        if self._heading is not None:
            self._heading_parts.append(data.strip())
        if data.strip():
            self.parts.append(data.strip())


def _html_text(raw: str) -> tuple[str, str | None]:
    parser = _HTMLText()
    parser.feed(raw)
    text = re.sub(r"\s+", " ", " ".join(parser.parts)).strip()
    return text, parser.title


def _pdf_text(path: Path) -> tuple[str, str | None] | None:
    try:
        from pypdf import PdfReader
    except ImportError:
        _missing_extra(path, "pdf")
        return None
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip(), None


def _docx_text(path: Path) -> tuple[str, str | None] | None:
    try:
        from docx import Document
        from docx.table import Table
    except ImportError:
        _missing_extra(path, "docx")
        return None
    document = Document(str(path))
    title = None
    parts = []
    def paragraphs(items):
        for item in items:
            if isinstance(item, Table):
                for row in item.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs
                        yield from paragraphs(cell.tables)
            else:
                yield item

    for para in paragraphs(document.iter_inner_content()):
        style = para.style.name if para.style is not None else ""
        text = para.text.strip()
        if text and title is None and style.startswith("Heading"):
            title = text
        if text:
            parts.append(text)
    return "\n".join(parts).strip(), title


def extract_file(path: Path) -> tuple[str, str | None] | None:
    """Return (text, title or None). None means the file was skipped."""
    suffix = path.suffix.lower()
    if suffix in (".md", ".markdown"):
        text = _read_text(path)
        return text, _markdown_title(text)
    if suffix == ".txt":
        return _read_text(path), None
    if suffix == ".vtt":
        return clean_vtt(_read_text(path)), None
    if suffix == ".srt":
        return clean_srt(_read_text(path)), None
    if suffix in (".html", ".htm"):
        return _html_text(_read_text(path))
    if suffix == ".pdf":
        return _pdf_text(path)
    if suffix == ".docx":
        return _docx_text(path)
    return None
